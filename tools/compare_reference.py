"""Acceptance test (brief section 8): compare a run's per-call and per-step ratings with Rory's reference report.
python -m tools.compare_reference "out/Call Scorecard Leanne Cooper, w-e 28 Aug 2026.json" [docs/reference-leanne-2026-08-28.docx]"""
import json
import sys

from docx import Document

from callscoring.config import STEPS

POINTS = {"Y": 1.0, "P": 0.5, "M": 0.0}


def reference(path):
    d = Document(path)
    grid = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == "Call" and len(t.columns) == 11)
    calls = [t for t in d.tables if [c.text.strip() for c in t.rows[0].cells][:3] == ["#", "Step", "Score"]]
    out = []
    for i, row in enumerate(grid.rows[1:]):
        cells = [c.text.strip() for c in row.cells]
        notes = [r.cells[3].text.strip() for r in calls[i].rows[1:]] if i < len(calls) else [""] * 9
        out.append({"customer": cells[0], "ratings": cells[1:10], "notes": notes})
    return out


def key(name):
    return (name or "").split()[0].lower()


def main():
    run = json.load(open(sys.argv[1], encoding="utf-8"))
    ref = reference(sys.argv[2] if len(sys.argv) > 2 else "docs/reference-leanne-2026-08-28.docx")
    ours = {key(c["customer"]): c for c in run["calls"]}
    print(f"Reference: {len(ref)} calls. This run: {len(run['calls'])} scored, {run['calls_made']} calls made, "
          f"{len(run['dropped'])} dropped, not scored {run['not_scored']}.\n")

    agree = total = 0
    step_ref = [0.0] * 9
    step_ours = [0.0] * 9
    for r in ref:
        c = ours.get(key(r["customer"]))
        if not c:
            print(f"MISSING  {r['customer']}: in the reference but not scored in this run")
            continue
        mine = [s["rating"] for s in c["score"]["steps"]]
        rp, op = sum(POINTS[x] for x in r["ratings"]), sum(POINTS[x] for x in mine)
        print(f"{r['customer']:<18} ref {''.join(r['ratings'])} {rp:.1f}   ours {''.join(mine)} {op:.1f}")
        for i in range(9):
            step_ref[i] += POINTS[r["ratings"][i]]
            step_ours[i] += POINTS[mine[i]]
            total += 1
            if mine[i] == r["ratings"][i]:
                agree += 1
                continue
            s = c["score"]["steps"][i]
            print(f"   step {i + 1} {STEPS[i]['short']}: ref {r['ratings'][i]}, ours {mine[i]}")
            print(f"      ref:  {r['notes'][i]}")
            print(f"      ours: {s['what_happened']}")
            if s["evidence"]:
                print(f"      evidence: \"{s['evidence']}\"")
    extra = [c["customer"] for k, c in ours.items() if k not in {key(r["customer"]) for r in ref}]
    if extra:
        print(f"\nEXTRA in this run (not in the reference): {', '.join(extra)}")
    n = len(ref)
    print(f"\nStep agreement {agree}/{total} ({100 * agree // max(total, 1)}%)")
    print(f"{'Step':<20}{'ref':>6}{'ours':>6}")
    for i, s in enumerate(STEPS):
        print(f"{s['short']:<20}{step_ref[i]:>6.1f}{step_ours[i]:>6.1f}")
    print(f"{'Total':<20}{sum(step_ref):>6.1f}{sum(step_ours):>6.1f}   (expected ref 31.0/{9 * n})")
    if run.get("week"):
        w = run["week"]
        print(f"Week: ours {w['overall_avg']:.1f}/9, {w['pct']}%, Must-Ask {w['must_avg']:.1f}  "
              f"(reference 3.4/9, 38%, Must-Ask 1.2)")


if __name__ == "__main__":
    main()
